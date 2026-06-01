import openpyxl
wb = openpyxl.load_workbook(r'd:\Solo Website\products_image_search_links.xlsx', data_only=True)
for s in wb.sheetnames:
    ws = wb[s]
    print(f'Sheet: {s}  rows={ws.max_row}  cols={ws.max_column}')
    headers = [c.value for c in ws[1]]
    print(f'  Headers: {headers}')
