"""Load product images from the candidates Excel report into Azure Blob + DB.

Reads:  d:\Solo Website\products_without_images_candidates.xlsx
        (sheet "Candidates Found" — all rows where Candidate Count > 0)

Writes: Azure Blob Storage (stsolowebsite / media / products)
        DB tables: media_assets, product_images

Run with firewall rule open to pg-qlyb5greec2io.postgres.database.azure.com:5432
"""

import os
import re
import sys
import time
import uuid
import json
import requests
import psycopg2
from pathlib import Path
from urllib.parse import unquote
from azure.storage.blob import BlobServiceClient, ContentSettings

# ── CONFIG ─────────────────────────────────────────────────────────

DB_CONFIG = {
    'host': 'pg-qlyb5greec2io.postgres.database.azure.com',
    'port': 5432,
    'dbname': 'solo_ecommerce',
    'user': 'soloadmin',
    'password': os.environ.get('PGPASSWORD') or (_ for _ in ()).throw(SystemExit('PGPASSWORD env var is required.')),
    'sslmode': 'require',
}

AZURE_CONN_STR = os.environ.get('AZURE_STORAGE_CONNECTION_STRING')
if not AZURE_CONN_STR:
    raise SystemExit('AZURE_STORAGE_CONNECTION_STRING env var is required.')

CONTAINER = 'media'
BLOB_FOLDER = 'products'

ROOT_DIR = Path(__file__).resolve().parents[1]
CANDIDATES_FILE = ROOT_DIR / 'products_without_images_candidates.xlsx'
PROGRESS_FILE = Path(__file__).resolve().parent / 'load_from_candidates_progress.json'

MIN_IMAGE_SIZE = 5000   # 5 KB — skip tiny placeholders

HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
                  '(KHTML, like Gecko) Chrome/131.0 Safari/537.36',
    'Accept': 'image/webp,image/apng,image/*,*/*;q=0.8',
    'Accept-Language': 'en-US,en;q=0.9',
}


# ── HELPERS ────────────────────────────────────────────────────────

def load_progress():
    if PROGRESS_FILE.exists():
        with open(PROGRESS_FILE, 'r') as f:
            return json.load(f)
    return {
        'done_product_ids': [],
        'stats': {
            'products_done': 0,
            'images_uploaded': 0,
            'images_skipped': 0,
            'errors': 0,
        },
    }


def save_progress(data):
    with open(PROGRESS_FILE, 'w') as f:
        json.dump(data, f, indent=2)


def download_image(url, session):
    """Download an image URL. Returns (bytes, content_type) or (None, None)."""
    try:
        resp = session.get(url, headers=HEADERS, timeout=30, allow_redirects=True)
        resp.raise_for_status()
        ct = resp.headers.get('Content-Type', 'image/jpeg').split(';')[0].strip()
        data = resp.content
        if len(data) < MIN_IMAGE_SIZE:
            return None, None
        return data, ct
    except Exception as e:
        return None, None


def upload_to_azure(container_client, image_data, content_type):
    """Upload image bytes to Azure Blob. Returns (blob_url, blob_name, filename, mime, file_uuid)."""
    file_uuid = str(uuid.uuid4())
    if 'webp' in content_type:
        mime, ext = 'image/webp', 'webp'
    elif 'png' in content_type:
        mime, ext = 'image/png', 'png'
    else:
        mime, ext = 'image/jpeg', 'jpg'

    filename = f'{file_uuid}.{ext}'
    blob_name = f'{BLOB_FOLDER}/{filename}'
    blob_client = container_client.get_blob_client(blob_name)
    blob_client.upload_blob(
        image_data,
        overwrite=True,
        content_settings=ContentSettings(
            content_type=mime,
            cache_control='public, max-age=31536000, immutable',
        ),
    )
    return blob_client.url, blob_name, filename, mime, file_uuid


def create_db_records(cur, file_uuid, filename, mime, size, alt_text,
                      product_id, display_order, is_primary):
    """Insert media_assets + product_images rows."""
    blob_key = f'{BLOB_FOLDER}/{filename}'
    cur.execute('''
        INSERT INTO media_assets (
            id, key, folder, filename, "mimeType", "sizeBytes",
            "altText", "isDeleted", "createdAt", "updatedAt"
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, false, NOW(), NOW())
    ''', (file_uuid, blob_key, BLOB_FOLDER, filename, mime, size, alt_text))

    cur.execute('''
        INSERT INTO product_images (
            product_id, media_asset_id, alt_text,
            display_order, is_primary, created_at, updated_at
        ) VALUES (%s, %s, %s, %s, %s, NOW(), NOW())
    ''', (product_id, file_uuid, alt_text, display_order, is_primary))


# ── EXCEL READER ──────────────────────────────────────────────────

def load_candidates():
    """
    Read the 'Candidates Found' sheet from the candidates Excel.
    Returns list of dicts with keys: product_id, sku, product_name, source, candidate_urls.
    """
    try:
        import openpyxl
    except ImportError:
        import subprocess
        subprocess.run([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'], check=True)
        import openpyxl

    wb = openpyxl.load_workbook(CANDIDATES_FILE, data_only=True)

    # Try "Candidates Found" first, fall back to "All Results"
    if 'Candidates Found' in wb.sheetnames:
        ws = wb['Candidates Found']
    else:
        ws = wb['All Results']

    # Column layout (1-indexed, matching the write_report() in the finder script):
    # 1=ID, 2=SKU, 3=Product Name, 4=Name(AR), 5=Slug, 6=Cat, 7=SubCat, 8=Brand,
    # 9=Price Incl, 10=Price Excl, 11=Stock, 12=Active, 13=Disc, 14=Short Desc,
    # 15=Created At, 16=Candidate Source, 17=Candidate Count, 18=Search URL,
    # 19=Product Page URL, 20=Candidate 1, 21=Candidate 2, 22=Candidate 3

    COL_ID    = 1
    COL_SKU   = 2
    COL_NAME  = 3
    COL_BRAND = 8
    COL_SRC   = 16
    COL_COUNT = 17
    COL_C1    = 20
    COL_C2    = 21
    COL_C3    = 22

    products = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[COL_ID - 1]:
            continue

        product_id  = str(row[COL_ID  - 1]).strip()
        sku         = str(row[COL_SKU  - 1]).strip() if row[COL_SKU - 1] else ''
        product_name = str(row[COL_NAME - 1]).strip() if row[COL_NAME - 1] else ''
        brand       = str(row[COL_BRAND - 1]).strip() if row[COL_BRAND - 1] else ''
        source      = str(row[COL_SRC  - 1]).strip() if row[COL_SRC - 1] else 'none'
        count       = int(row[COL_COUNT - 1]) if row[COL_COUNT - 1] else 0

        if count == 0 or source == 'none':
            continue

        candidate_urls = []
        for col in [COL_C1, COL_C2, COL_C3]:
            val = row[col - 1]
            if val and str(val).strip().startswith('http'):
                candidate_urls.append(str(val).strip())

        if not candidate_urls:
            continue

        products.append({
            'product_id':   product_id,
            'sku':          sku,
            'product_name': product_name,
            'brand':        brand,
            'source':       source,
            'candidate_urls': candidate_urls,
        })

    wb.close()
    return products


# ── MAIN ──────────────────────────────────────────────────────────

def main():
    print('=' * 65, flush=True)
    print('LOAD IMAGES FROM CANDIDATES  (Excel → Azure Blob + DB)', flush=True)
    print('=' * 65, flush=True)

    print(f'\n[1/4] Reading candidates from: {CANDIDATES_FILE.name}', flush=True)
    if not CANDIDATES_FILE.exists():
        print(f'  ERROR: File not found: {CANDIDATES_FILE}', flush=True)
        sys.exit(1)

    products = load_candidates()
    print(f'  Loaded {len(products)} products with candidates', flush=True)

    # Count by source
    from collections import Counter
    src_counts = Counter(p['source'] for p in products)
    for src, cnt in sorted(src_counts.items()):
        print(f'    {src}: {cnt}', flush=True)

    if not products:
        print('  Nothing to do.', flush=True)
        return

    # Load progress
    progress = load_progress()
    done_ids = set(progress['done_product_ids'])
    stats    = progress['stats']
    pending  = [p for p in products if p['product_id'] not in done_ids]
    print(f'  Already done: {len(done_ids)}  |  Pending: {len(pending)}', flush=True)

    if not pending:
        print('\n  All products already processed!', flush=True)
        return

    print('\n[2/4] Connecting to Azure Blob Storage...', flush=True)
    blob_service     = BlobServiceClient.from_connection_string(AZURE_CONN_STR)
    container_client = blob_service.get_container_client(CONTAINER)
    print('  Connected.', flush=True)

    print('\n[3/4] Connecting to PostgreSQL...', flush=True)
    conn = psycopg2.connect(**DB_CONFIG)
    cur  = conn.cursor()
    print('  Connected.', flush=True)

    print('\n[4/4] Processing products...\n', flush=True)
    session = requests.Session()
    total   = len(pending)

    for idx, product in enumerate(pending, start=1):
        pid   = product['product_id']
        sku   = product['sku']
        name  = product['product_name']
        src   = product['source']
        urls  = product['candidate_urls']

        # Safety: verify this product really has no images in DB (avoid duplicates)
        cur.execute(
            'SELECT COUNT(*) FROM product_images WHERE product_id = %s', (pid,)
        )
        existing_count = cur.fetchone()[0]
        if existing_count > 0:
            print(f'  SKIP {sku} — already has {existing_count} image(s) in DB', flush=True)
            done_ids.add(pid)
            stats['products_done'] += 1
            continue

        # Download & upload each candidate URL
        product_img_count = 0
        for display_order, img_url in enumerate(urls):
            try:
                cur.execute('SAVEPOINT img_sp')

                image_data, content_type = download_image(img_url, session)
                if not image_data:
                    cur.execute('RELEASE SAVEPOINT img_sp')
                    stats['images_skipped'] += 1
                    continue

                _, _, filename, mime, file_uuid = upload_to_azure(
                    container_client, image_data, content_type
                )

                is_primary = (product_img_count == 0)
                create_db_records(
                    cur, file_uuid, filename, mime,
                    len(image_data), name,
                    pid, display_order, is_primary
                )

                cur.execute('RELEASE SAVEPOINT img_sp')
                product_img_count += 1
                stats['images_uploaded'] += 1
                time.sleep(0.15)

            except Exception as e:
                print(f'    ERR {sku} img#{display_order}: {e}', flush=True)
                try:
                    cur.execute('ROLLBACK TO SAVEPOINT img_sp')
                except Exception:
                    pass
                stats['errors'] += 1

        conn.commit()
        done_ids.add(pid)
        stats['products_done'] += 1

        status = f'{product_img_count} img(s) [{src}]'
        if product_img_count == 0:
            status = 'NO IMAGES DOWNLOADED'
        print(f'  [{idx:>3}/{total}] {sku}  {name[:40]}  →  {status}', flush=True)

        # Save progress every 20 products
        if idx % 20 == 0:
            save_progress({'done_product_ids': list(done_ids), 'stats': stats})

    # Final save
    save_progress({'done_product_ids': list(done_ids), 'stats': stats})
    conn.close()

    print('\n' + '=' * 65, flush=True)
    print('DONE', flush=True)
    print(f"  Products processed : {stats['products_done']}", flush=True)
    print(f"  Images uploaded    : {stats['images_uploaded']}", flush=True)
    print(f"  Images skipped     : {stats['images_skipped']}", flush=True)
    print(f"  Errors             : {stats['errors']}", flush=True)
    print('=' * 65, flush=True)


if __name__ == '__main__':
    main()
